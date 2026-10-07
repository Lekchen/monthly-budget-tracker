import json
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
print("==============================")
print(" MONTHLY BUDGET TRACKER")
print("==============================")

expenses = []
monthly_income = 0
budgets = {} 
try:
    with open("budget_data.json", "r") as file:
        data = json.load(file)

    monthly_income = data["monthly_income"]
    budgets = data["budgets"]
    expenses = data["expenses"]

    print("Saved data loaded successfully!")

except FileNotFoundError:
    print("No saved data found.")
while True:
    print()
    print("==============================")
    print("            MENU")
    print("==============================")
    print("1. Add Expense")
    print("2. View Expenses")
    print("3. Set Monthly Income")
    print("4. Set Category Budget")
    print("5. View Budget Summary")
    print("6. Save Data")
    print("7. Generate Excel Report")
    print("8. Exit")

    choice = input("Enter your choice: ")
    if choice == "1":
        while True:
            date = input("Enter date (YYYY-MM-DD): ")

            try:
                datetime.strptime(date, "%Y-%m-%d")
                break
            except ValueError:
                print("Invalid date. Please use YYYY-MM-DD.")
        description = input("Enter description: ")
        category = input("Enter category: ").lower()
        while True:
            try:
                amount = float(input("Enter amount: $"))
                if amount <= 0:
                    print("Amount must be greater than $0.")
                else:
                    break
            except ValueError:
                print("Invalid amount. Please enter a number.")
        expense = {
            "date": date,
            "description": description,
            "category": category,
            "amount": amount
        }
        expenses.append(expense)
        print("Expense added successfully!")
    elif choice == "2":
        if len(expenses) == 0:
            print("No expenses found.")
        else:
            print("==============================")
            print("        YOUR EXPENSES")
            print("==============================")
            total = 0
            for expense in expenses:
                print(
                    expense.get("date", "No date"),
                    "|",
                    expense["description"],
                    "|",
                    expense["category"],
                    "|",
                    f'${expense["amount"]:.2f}'
                )
                total += expense["amount"]

            print("==============================")
            print("Total spent:", f"${total:.2f}")
    elif choice == "3":
        while True:
            try:
                monthly_income = float(input("Enter your monthly income: $"))

                if monthly_income <= 0:
                    print("Income must be greater than $0.")
                else:
                    break

            except ValueError:
                print("Invalid income. Please enter a number.")

        print("Monthly income set successfully!")
    elif choice == "4":
        category = input("Enter category: ").lower()

        while True:
            try:
                budget_amount = float(input("Enter budget amount: $"))

                if budget_amount <= 0:
                    print("Budget must be greater than $0.")
                else:
                    break

            except ValueError:
                print("Invalid budget. Please enter a number.")

        budgets[category] = budget_amount
        print("Budget set successfully!")
    elif choice == "5":
        print("========================")
        print("      BUDGET SUMMARY    ")
        print("========================")

        print("Monthly Income:", f"${monthly_income:.2f}")
        total_budget = 0
        for category in budgets:
            spent = 0
            for expense in expenses:
                if expense["category"] == category:
                    spent += expense["amount"]
            remaining = budgets[category] - spent
            print(
                category,
                "| Budget:", f"${budgets[category]:.2f}",
                "| Spent:", f"${spent:.2f}",
                "| Remaining:", f"${remaining:.2f}"
)
            total_budget += budgets[category]
        print("Total Budget:", f"${total_budget:.2f}")
    elif choice == "6":
        data = {
            "monthly_income": monthly_income,
            "budgets": budgets,
            "expenses": expenses
        }
        with open("budget_data.json", "w") as file:
            json.dump(data, file, indent=4)
        print("Data saved successfully!")
    elif choice == "7":
        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "Expenses"

        sheet.append(["Date", "Description", "Category", "Amount"])

        for expense in expenses:
            sheet.append([
                expense.get("date", "No date"),
                expense["description"],
                expense["category"],
                expense["amount"]
            ])

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center")

        for row in range(2, sheet.max_row + 1):
            sheet.cell(row=row, column=4).number_format = "$0.00"

        sheet.column_dimensions["A"].width = 15
        sheet.column_dimensions["B"].width = 25
        sheet.column_dimensions["C"].width = 20
        sheet.column_dimensions["D"].width = 15

        budget_sheet = workbook.create_sheet("Budget Summary")

        budget_sheet.append([
            "Category",
            "Budget",
            "Spent",
            "Remaining"
        ])

        total_budget = 0

        total_spent = sum(
            expense["amount"] for expense in expenses
        )

        for category in budgets:
            spent = 0

            for expense in expenses:
                if expense["category"] == category:
                    spent += expense["amount"]

            remaining = budgets[category] - spent

            budget_sheet.append([
                category,
                budgets[category],
                spent,
                remaining
            ])

            total_budget += budgets[category]


        remaining_budget = total_budget - total_spent
        savings = monthly_income - total_spent

        budget_sheet.append([])

        budget_sheet.append(["Monthly Income", monthly_income])
        budget_sheet.append(["Total Budget", total_budget])
        budget_sheet.append(["Total Spent", total_spent])
        budget_sheet.append(["Remaining Budget", remaining_budget])
        budget_sheet.append(["Savings", savings])

        for cell in budget_sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.alignment = Alignment(horizontal="center")

        for row in range(2, len(budgets) + 2):
            budget_sheet.cell(
                row=row,
                column=2
            ).number_format = "$0.00"

            budget_sheet.cell(
                row=row,
                column=3
            ).number_format = "$0.00"

            budget_sheet.cell(
                row=row,
                column=4
            ).number_format = "$0.00"

        summary_start_row = len(budgets) + 3

        for row in range(
            summary_start_row,
            summary_start_row + 5
        ):
            budget_sheet.cell(
                row=row,
                column=1
            ).font = Font(bold=True)

            budget_sheet.cell(
                row=row,
                column=2
            ).number_format = "$0.00"

        budget_sheet.column_dimensions["A"].width = 25
        budget_sheet.column_dimensions["B"].width = 18
        budget_sheet.column_dimensions["C"].width = 18
        budget_sheet.column_dimensions["D"].width = 18

        chart = BarChart()

        chart.title = "Budget vs Spent by Category"
        chart.y_axis.title = "Amount ($)"
        chart.x_axis.title = "Category"

        data = Reference(
            budget_sheet,
            min_col=2,
            max_col=3,
            min_row=1,
            max_row=len(budgets) + 1
        )

        categories = Reference(
            budget_sheet,
            min_col=1,
            min_row=2,
            max_row=len(budgets) + 1
        )

        chart.add_data(
            data,
            titles_from_data=True
        )

        chart.set_categories(categories)

        chart.height = 8
        chart.width = 14

        budget_sheet.add_chart(
            chart,
            "F2"
        )

        pie_chart = PieChart()

        pie_chart.title = "Spending by Category"

        pie_data = Reference(
            budget_sheet,
            min_col=3,
            min_row=1,
            max_row=len(budgets) + 1
        )
        pie_categories = Reference(
            budget_sheet,
            min_col=1,
            min_row=2,
            max_row=len(budgets) + 1
        )

        pie_chart.add_data(
            pie_data,
            titles_from_data=True
        )

        pie_chart.set_categories(
            pie_categories
        )

        pie_chart.height = 8
        pie_chart.width = 11

        pie_chart.layout = Layout(
            manualLayout=ManualLayout(
                x=0.00,
                y=0.00,
                w=0.60,
                h=0.60
            )
        )
        pie_chart.legend.position = "b"

        budget_sheet.add_chart(
            pie_chart,
            "F18"
        )
        workbook.save("monthly_budget_report.xlsx")
        print("Excel report generated successfully!")

    elif choice == "8":
        print("Goodbye!")
        break
    else:
        print("Invalid choice")